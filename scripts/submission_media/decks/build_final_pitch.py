"""Build the owner-directed six-slide pitch without changing frozen RC media.

Run with the bundled Python, then export_final_pitch.ps1 for native PDF/previews.
All diagrams and slide text are editable PowerPoint objects.
"""
from pathlib import Path
from html import escape
import hashlib
import json
import shutil
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.util import Inches, Pt
from PIL import Image, ImageOps, ImageDraw

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / 'docs/submission/media/decks/final'
OUT.mkdir(parents=True, exist_ok=True)
NAME = 'savia-final-pitch'
INK, MINT, WHITE, MUTED, PANEL, GOLD = '12332D', 'D9EDB7', 'F8F8F2', 'C6D3CA', '214A3F', 'E9C977'
REPO = 'https://github.com/mario-andreschak/factored-hackathon-2026-mcg'
CAP_PIN = '95e2b9d62685f2669473258b88935a43dbe1bf75'
SOURCE_PIN = '7de0bbd670c25e4ad83476c6d7b70c45cc58a137'
APP_PIN = 'f3c57b26d07c1e96b6e61a25befcbbc5fd51a17f'
SOURCES = {
    'capacity': f'{REPO}/blob/{CAP_PIN}/docs/submission/measurements/INFRASTRUCTURE_CAPACITY.md',
    'capacity_receipt': f'{REPO}/blob/{CAP_PIN}/docs/submission/measurements/infrastructure-capacity.json',
    'release': f'{REPO}/blob/843da50698b559dc748bea4505d1d93483ff29b3/docs/submission/RELEASE_CANDIDATE.md',
    'customer': f'{REPO}/blob/{SOURCE_PIN}/docs/submission/CUSTOMER_JOURNEY.md',
    'connector': f'{REPO}/blob/{SOURCE_PIN}/docs/submission/assistant/FLEET_CONNECTOR.md',
    'engine': f'{REPO}/blob/{SOURCE_PIN}/docs/DISPUTE_IMPLEMENTATION.md',
    'core_checks': f'{REPO}/blob/{SOURCE_PIN}/docs/submission/measurements/core-engine-final/README.md',
    'luna': f'{REPO}/blob/{SOURCE_PIN}/docs/submission/measurements/luna-100/README.md',
    'card': f'{REPO}/blob/{SOURCE_PIN}/docs/submission/measurements/CARD_BLOCK_VERIFICATION.md',
    'card_live': f'{REPO}/blob/{SOURCE_PIN}/docs/submission/measurements/card-block-live/README.md',
    'card_app': f'{REPO}/blob/{APP_PIN}/frontend/server/app.py',
    'operating_decisions': f'{REPO}/blob/{SOURCE_PIN}/docs/review/OPERATING_DECISIONS.md',
    'comparison': f'{REPO}/blob/{SOURCE_PIN}/docs/submission/ELEVENLABS_COMPARISON.md',
    'boundary': f'{REPO}/blob/{SOURCE_PIN}/docs/FLUJO_PRODUCT_BOUNDARY.md',
    'landscape': f'{REPO}/blob/{SOURCE_PIN}/docs/architecture/system-landscape.md',
    'recovery': f'{REPO}/blob/{SOURCE_PIN}/docs/architecture/rc-runtime-reference.md',
    'models': 'https://github.com/mario-andreschak/FLUJO/blob/main/docs/features/models/connecting.md',
    'eleven': 'https://elevenlabs.io/blog/eleven-v4-turbo-in-elevenagents',
    'pitch_style': 'https://factored-hackathon.slack.com/archives/C0BUZCY0TUY/p1791068906645819',
    'english': 'https://factored-hackathon.slack.com/archives/C0BUZCY0TUY/p1791153414484629',
    'submission': 'https://factored-hackathon.slack.com/archives/C0BU6V61273/p1791211945034199',
    'voice_clarification': 'https://factored-hackathon.slack.com/archives/C0BU6V61273/p1791228494554899',
}

SLIDES = [
    dict(title='Ask once.\nSavia follows through.', kicker='SAVIA / FACTORED AI & DATA HACKATHON 2026',
         speak="I don't recognize this transaction. That can start a problem that takes more than one conversation to solve. Meet Savia: one assistant that keeps your case and brings back a useful next step. The charge stays visible, so the answer stays connected to the problem. Ask once. Savia follows through while you get on with your day.",
         evidence='Implemented core: deterministic R0–R18 dispute flow, owned facts, ordered policy, explicit consent, verified receipts, recovery and handoff. Recorded prototype: integrated voice, grounded explanation, durable inquiries and two completed model reviewers. Immediate fictional card protection has separate local and public API acceptance. Voice and larger collaborative fleets extend the trusted core; neither transfers its acceptance to a whole customer journey.', sources=['engine','core_checks','customer','card_live']),
    dict(title='Your problem keeps moving.\nYou get on with your day.', kicker='01 / THE CUSTOMER EXPERIENCE',
         speak='Select the unfamiliar charge. Savia explains the checked merchant, amount and status. Ask its specialists to compare evidence; their saved suggestions wait when you return. The recorded customer path includes a grounded answer, two completed reviewers, and saved recommendations after a new chat. Confirm once. Savia blocks the owned card, writes a durable status, and rereads an independent receipt. Asking never writes. Foreign cards, missing consent, and tampered receipts are refused.',
         evidence='Actual fictional Spanish UI is shown with an English summary. Recorded: grounded answer, two completed model reviewers, saved suggestions surviving new chat/reload and helpful informational closure. Complete native voice playback has receipts; a separate simulated receipt check repeats after 30 real minutes. New protection evidence consists of two public authenticated API journeys over fictional cards, not a new GUI or voice recording. Preparation leaves status unchanged; explicit consent precedes verified block readback, and a fresh login recovers it. Real human assignment, real-bank block/refund, push/email delivery and multi-day completion remain unverified. The story rail joins separately qualified capabilities; it is not a recorded production lifecycle.', sources=['customer','release','card','card_live']),
    dict(title='Voice is the beginning.\nFollow-through is the product.', kicker='02 / THE COMMERCIAL REFERENCE',
         speak='A polished voice is not a case. ElevenLabs is a commercial voice platform. Savia is the bank-controlled service: owned facts, deterministic policy, consent, receipt, follow-up, swarm investigation. We built that in ten days with two people.',
         evidence='ElevenLabs official September 28 article, updated October 5 and checked October 5, reports expressive voice, about 100 ms median model inference, 90+ languages, a financial-services dispute showcase and a combined transcription/turn-taking/TTS stack. These are vendor-reported capabilities; no matched performance benchmark is available. The comparison describes Savia’s product emphasis and bank-controlled policy/action boundary. The two-person, ten-day delivery statement is owner-supplied; historical implementation credits for Gloria Yanta Salc, Carlos Diaz and Mario Andreschak remain preserved.', sources=['eleven','comparison','boundary']),
    dict(title='One customer voice.\nTen teams behind it.', kicker='03 / THE SWARM DESIGN',
         speak='One customer voice. Ten teams behind it. 300/300 concurrent FLUJO. 18 sandboxes live together. The customer path we recorded is two reviewers; the architecture is ready to scale.',
         evidence='Exact 100-agent customer completion has not yet been demonstrated; one later attempt ended on a disabled provider workspace before delegation. Diagram counting: ten team Machines × (one lead + nine child specialists) = 100 team conversations; the root coordinator is separate. The 300/300 foundation result is the repeat neutral reference-code workload with 300 concurrent client submissions queued through FLUJO, 1k context, Qwen3.8-27B FP8 and vLLM on a rented H100; it does not mean 300 simultaneous GPU generations or 300 customer investigations. Queue-inclusive p50/p95 were 273.40/441.60 seconds. First wording yielded 217/300 correct; direct inference yielded 399/400. Historical collaboration had 18 leaf sandboxes live together, 17 completed leaf team runs and six matching output hashes. Six of 24 leaf slots failed provisioning, and operator steering was required. The larger recovered run enrolled 96 child workers over time. These preserve the tested foundation and qualify the scale design; they are separate from the successful recorded two-reviewer customer path and the local Luna workload.', sources=['connector','capacity','capacity_receipt','customer','recovery']),
    dict(title='Evidence you can inspect.\nChoices a bank can control.', kicker='04 / QUALIFIED PROOF & DATA DECISIONS',
         speak='Policy, not the model, authorizes the card action. Savia and Banking MCP own facts, nineteen ordered rules, consent and receipts. FLUJO remains the generic orchestration foundation. One thousand seventy-eight offline core checks exercise that boundary. One hundred local Luna fixture decisions matched their expected results. Two live Spanish and Portuguese API journeys verified card protection. The data audit taught us to ground new inquiries in owned transactions. The bank controls the stack and the next step.',
         evidence='Fresh core qualification: 1,078 passing checks, zero errors/failures/skips and no provider calls. It executed 14 suites; 39 core/domain source files independently match the deployed source. This does not establish all-repository CI or current end-to-end customer acceptance. Luna: 100 completed actual local subscription requests, 50 ES/50 PT, 10 distinct request phrases with varied fictional facts, exact frozen oracle and bounded no-action checks. Independent deterministic baseline also 100/100; no quality improvement, 100 collaborative agents, deployed GUI, bank writes or simultaneous GPU-generation claim. Public API protection: two actual ES/PT authenticated journeys on application f3c57b26, consent before a fictional block, verified receipt readback/retry and status recovered on fresh login; no new GUI or voice acceptance. Data-to-product decision: published complaint-product links violate customer ownership, so new inquiries use owned reads. Keep ownership and no-write rules deterministic. Savia/Banking MCP own domain data/policy/state; FLUJO main keeps generic chat/flow/tool/MCP interfaces. Source-documented local-model/deployment options are not a measured all-local Savia rollout. Historical queued reference-code capacity, including 300/300 repeat workload, remains separate extension evidence with failure and timing disclosures.', sources=['core_checks','luna','card_live','card_app','operating_decisions','boundary','landscape','models','capacity','capacity_receipt']),
    dict(title='Make “I’ll look into it”\na service customers can feel.', kicker='05 / THE BANK PILOT',
         speak='For customers, the ambition is less chasing and less repeating. For the bank, it is clearer evidence and a better prepared handoff. Start with transaction disputes, then reuse the same modular service for other customer problems. We are seeking a bank partner for a bounded pilot: measure repeat contacts, helpful answers, handoff quality and cost per case. Measure the service on real customer needs, with a small rollout and a clear review gate. Ask once. Savia follows through.',
         evidence='Pilot invitation, not measured ROI. Live fictional demo link is a review entry point, not a freshly verified end-to-end customer acceptance. Entry code SAVIA-2026. Production bank actions and live human assignment are not claimed. Founder/team credits preserve Gloria Yanta Salc (prompt flow / decision design), Carlos Diaz (data pipeline / lookup), Mario Andreschak (integration / orchestration).', sources=['customer','release']),
]

def color(c): return RGBColor.from_string(c)

def box(s, x,y,w,h, fill=PANEL, line=None, rounded=True):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=color(fill)
    sh.line.fill.background() if not line else None
    if line: sh.line.color.rgb=color(line)
    if rounded: sh.adjustments[0]=.12
    return sh

def txt(s,x,y,w,h,value,size=20,c=WHITE,bold=False,link=None):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.01); tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(2)
        r=p.add_run(); r.text=line; r.font.name='Aptos'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color(c)
    if link: sh.click_action.hyperlink.address=link
    return sh

def circle(s,x,y,d,c):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)); sh.fill.solid(); sh.fill.fore_color.rgb=color(c); sh.line.fill.background(); return sh

def line(s,x1,y1,x2,y2,c=MINT,width=1.5):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); sh.line.color.rgb=color(c); sh.line.width=Pt(width)

def eyes(s,x,y,d=1.2):
    for dx,dy in [(0,.07),(d*.98,0)]:
        circle(s,x+dx,y+dy,d,WHITE); circle(s,x+dx+d*.43,y+dy+d*.39,d*.43,INK)

def base(prs,i):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=color(INK)
    box(s,.64,.56,.50,.07,MINT,rounded=False)
    txt(s,1.32,.41,13,.35,SLIDES[i]['kicker'],11,MINT,True)
    txt(s,.72,8.50,12,.25,'SAVIA  /  ASK ONCE. SAVIA FOLLOWS THROUGH.',9,MUTED)
    txt(s,14.22,8.48,1,.28,f'{i+1:02} / 06',10,MINT,True)
    if i>0: txt(s,.72,1.07,14.8,1.58,SLIDES[i]['title'],38,WHITE,True)
    s.notes_slide.notes_text_frame.text='SAY (participant voice)\n'+SLIDES[i]['speak']+'\n\nIF ASKED\n'+SLIDES[i]['evidence']+'\n\nSOURCES\n'+'\n'.join(SOURCES[k] for k in SLIDES[i]['sources'])
    return s

def card(s,x,y,w,h,kicker,title,body,accent=MINT):
    box(s,x,y,w,h)
    txt(s,x+.22,y+.18,w-.44,.32,kicker,10,accent,True)
    txt(s,x+.22,y+.64,w-.44,.69,title,23,WHITE,True)
    txt(s,x+.22,y+1.42,w-.44,h-1.56,body,16,MUTED)

def build():
    prs=Presentation(); prs.slide_width=Inches(16); prs.slide_height=Inches(9)
    prs.core_properties.title='Savia — a decision engine with saved follow-through'; prs.core_properties.author='Savia team'; prs.core_properties.subject='Final six-slide bank-pilot product pitch / October 5, 2026'
    s=base(prs,0)
    txt(s,.75,1.48,9.3,2.15,SLIDES[0]['title'],55,WHITE,True)
    txt(s,.79,4.04,8.5,1.35,'One voice for you.\nA case that keeps its context.',27,MINT)
    eyes(s,11.07,1.74,1.59)
    for j in range(16): box(s,10.25+j*.27,4.40-((j%5)*.14),.08,.42+(j%5)*.28,MINT,rounded=False)
    box(s,.79,6.07,14.40,1.53)
    txt(s,1.06,6.31,9.5,.60,'“I don’t recognize this transaction.”',28,WHITE,True)
    txt(s,1.08,7.02,9.8,.28,'OUR FIRST USE CASE: TRANSACTION DISPUTES',11,MINT,True)
    txt(s,11.05,6.41,3.85,.68,'Checked facts.\nA saved next step.',16,MUTED)

    s=base(prs,1)
    stages=[('Check','Understand your selected charge.'),('Protect','Confirm once. Block your owned card.'),('Explore','Specialists compare the evidence.'),('Return','Read your saved answer and next step.')]
    for j,(title,body) in enumerate(stages):
        x=.78+j*2.64
        circle(s,x,3.04,.36,MINT); txt(s,x+.49,3.02,1.9,.36,f'{j+1:02}',13,MINT,True)
        if j<3: line(s,x+.39,3.22,x+2.54,3.22)
        txt(s,x,3.75,2.48,.46,title,27,WHITE,True)
        txt(s,x,4.35,2.42,1.15,body,18,MUTED)
    box(s,.78,6.07,10.35,1.57)
    txt(s,1.02,6.28,9.8,.28,'CARD PROTECTION / COMPLETED ACTION',10,MINT,True)
    txt(s,1.02,6.77,9.75,.98,'Confirm once. Savia blocks the owned card, writes a durable status, and rereads an independent receipt. Asking never writes. Foreign cards, missing consent, and tampered receipts are refused.',16,WHITE)
    shot=OUT/'savia-integrated-closure-crop.png'
    shutil.copyfile(ROOT/'docs/submission/media/decks/savia-integrated-closure-crop.png',shot)
    s.shapes.add_picture(str(shot), Inches(11.78), Inches(2.95), width=Inches(3.34), height=Inches(3.78))
    txt(s,11.77,6.91,3.38,.93,'English summary: “The explanation helped.” Saved suggestions remain available.',15,MINT)
    txt(s,.79,7.96,14.45,.31,'Demo ledger. Same admission rules a production host would use.',11,MUTED)

    s=base(prs,2)
    card(s,.79,3.03,6.93,3.90,'ELEVENAGENTS / COMMERCIAL VOICE REFERENCE','A compelling conversation','Expressive, responsive voice\nFinancial-services dispute showcase\nIntegrated speech and turn-taking',GOLD)
    card(s,8.01,3.03,7.18,3.90,'SAVIA / THE CONTINUING CASE','A case that keeps moving','Checked facts and ordered decisions\nExplicit consent and saved receipts\nSaved context and specialist perspectives')
    txt(s,.82,7.29,14.2,.54,'A familiar voice. A persistent case. A clear next step.',25,MINT,True)
    txt(s,.82,7.96,14.2,.27,'Two people. Ten days. A service the bank can inspect and control.',12,MUTED)
    txt(s,.82,8.22,13,.23,'Reference: ElevenLabs — Eleven v4 Turbo in ElevenAgents',9,MUTED,link=SOURCES['eleven'])

    s=base(prs,3)
    txt(s,.81,2.98,4.1,1.90,'100',102,MINT,True)
    txt(s,.87,5.04,4.0,.73,'team conversations',24,WHITE,True)
    txt(s,.87,5.93,4.1,.95,'SWARM DESIGN\n10 teams × (1 lead + 9 specialists)',16,MUTED)
    for j in range(10):
        col,row=j%5,j//5; x=5.57+col*1.93; y=3.01+row*1.59
        box(s,x,y,1.73,1.38)
        txt(s,x+.14,y+.13,1.45,.23,f'TEAM {j+1:02}',9,MUTED,True)
        for n in range(10): circle(s,x+.20+(n%5)*.27,y+.54+(n//5)*.33,.17,GOLD if n==0 else MINT)
    txt(s,5.59,6.25,9.55,.29,'SHARED EVIDENCE  /  SPECIALIST COLLABORATION  /  INDEPENDENT REVIEW',10,MINT,True)
    for j,(title,body) in enumerate([
        ('300 / 300','Concurrent FLUJO reference returns'),
        ('18 live together','Sandboxes in the tested foundation'),
        ('Two completed reviewers','The customer path we recorded'),
    ]):
        x=.80+j*4.84; box(s,x,7.02,4.60,.95)
        txt(s,x+.18,7.14,4.2,.32,title,17,WHITE,True)
        txt(s,x+.18,7.58,4.2,.24,body,11,MUTED)
    txt(s,.82,8.13,14.2,.25,'One customer voice. Ten teams behind it. The architecture is ready to scale.',12,MINT)

    s=base(prs,4)
    txt(s,.82,2.90,14.3,.40,'Policy, not the model, authorizes the card action.',22,MINT,True)
    nodes=[('OWNED FACTS','Customer + charge'),('ORDERED POLICY','R0–R18 + consent'),('VERIFIED RECEIPT','Readback + recovery'),('SAVED FOLLOW-UP','Context + handoff')]
    for j,(k,title) in enumerate(nodes):
        x=.80+j*3.65; box(s,x,3.52,3.38,.94)
        txt(s,x+.16,3.67,3.05,.23,k,9,MINT,True)
        txt(s,x+.16,4.05,3.05,.29,title,16,WHITE,True)
        if j<3: txt(s,x+3.39,3.89,.25,.32,'→',14,MINT,True)
    for j,(stat,label,detail) in enumerate([
        ('1,078','offline core checks','R0–R18, ownership, consent,\nreceipts and recovery'),
        ('100 / 100','local Luna fixture decisions','50 ES + 50 PT\nGrounded expected decisions'),
        ('2 journeys','live public API · card protection','Spanish + Portuguese: consent,\nverified block and status on new login'),
    ]):
        x=.80+j*4.84; box(s,x,4.73,4.60,1.92)
        txt(s,x+.19,4.89,4.21,.58,stat,32,MINT,True)
        txt(s,x+.19,5.56,4.21,.28,label,13,WHITE,True)
        txt(s,x+.19,5.99,4.21,.47,detail,11,MUTED)
    box(s,.80,6.92,14.28,1.00)
    txt(s,1.00,7.06,13.85,.21,'DATA → PRODUCT DECISION',9,MINT,True)
    txt(s,1.00,7.40,13.85,.28,'Historical links crossed customer ownership. New inquiries use owned transaction reads.',17,WHITE,True)
    txt(s,.82,8.13,14.4,.23,'Savia + Banking MCP own the banking boundary. FLUJO keeps its generic interfaces; voice and models remain replaceable.',10,MUTED)

    s=base(prs,5)
    for j,(k,title,body) in enumerate([
        ('FOR THE CUSTOMER','Less chasing.','One guide keeps the context\nand brings the next step back.'),
        ('FOR THE BANK','Better prepared support.','Reviewed evidence and a clear\npacket for the next handoff.'),
        ('THE NEXT STEP','A bounded bank pilot.','Start with transaction disputes.\nMeasure value before scaling.')]):
        card(s,.80+j*4.84,3.03,4.60,2.94,k,title,body)
    box(s,.80,6.24,14.28,1.57,MINT)
    txt(s,1.05,6.44,13.74,.47,'Ask once. Savia follows through.',29,INK,True)
    txt(s,1.07,7.09,13.66,.33,'Pilot measures: repeat contacts · helpful answers · handoff quality · cost per case',14,INK)
    txt(s,.84,8.00,7.10,.24,'Try the fictional demo: savia-rc-2026.fly.dev  /  SAVIA-2026',10,MINT,link='https://savia-rc-2026.fly.dev')
    txt(s,8.15,8.00,7,.23,'Gloria Yanta Salc · Carlos Diaz · Mario Andreschak',9,MUTED)
    prs.save(OUT/f'{NAME}.pptx')
    (OUT/f'{NAME}-speaker-notes.json').write_text(json.dumps([{'slide':i+1,**s} for i,s in enumerate(SLIDES)],ensure_ascii=False,indent=2),encoding='utf-8')
    notes=['# Savia final pitch — narration and evidence', '', 'Six slides. Narration is approximately three minutes at a comfortable pace; record in a participant’s actual voice.', '']
    for i,spec in enumerate(SLIDES,1):
        notes += [f"## {i}. {spec['title'].replace(chr(10),' ')}",'','**Say:** '+spec['speak'],'','**If asked:** '+spec['evidence'],'','Sources: '+ ' · '.join(f'[{k}]({SOURCES[k]})' for k in spec['sources']),'']
    (OUT/f'{NAME}-speaker-notes.md').write_text('\n'.join(notes),encoding='utf-8')
    provenance={'version':'savia-final-pitch/v2','date':'2026-10-05','slide_count':6,'narration_words':sum(len(s['speak'].split()) for s in SLIDES),'status':'owner-review-successor-to-frozen-RC-decks','source_pin':SOURCE_PIN,'live_application_pin':APP_PIN,'sources':SOURCES,'claims':SLIDES,'organizer_style':'60% product / 40% technical; English; 4–6 slides','audio_guidance':'AI audio allowed with lower presentation score; actual participant voice recommended. Captured assistant speech has no confirmed exemption.','design':'Slides 1–3 and 6 product; slides 4–5 technical. Customer-led pitch; tested foundation and 100-conversation swarm design on slide 4; deterministic engine and scoped modern proof on slide 5. One message per slide; editable diagrams; English summary of actual Spanish evidence. Spoken product story is separated from If asked qualification details.','preserved_frozen_artifacts':{}}
    for rel in ['docs/submission/media/decks/savia-pitch-deck.pptx','docs/submission/media/decks/savia-pitch-deck.pdf','docs/submission/media/decks/savia-submission-deck.pptx','docs/submission/media/decks/savia-submission-deck.pdf','docs/submission/media/video/savia-submission.mp4','docs/submission/media/media-freeze-receipt.json']:
        provenance['preserved_frozen_artifacts'][rel]=hashlib.sha256((ROOT/rel).read_bytes()).hexdigest()
    (OUT/f'{NAME}-sources.json').write_text(json.dumps(provenance,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'Built {len(prs.slides)} editable slides at {OUT}; narration {sum(len(s["speak"].split()) for s in SLIDES)} words.')

def finish():
    previews=sorted((OUT/'slides').glob('slide-*.png'))
    if len(previews)!=6: raise RuntimeError('Export six slide PNGs first.')
    canvas=Image.new('RGB',(1640,1450),'#EAECE5'); draw=ImageDraw.Draw(canvas)
    for i,p in enumerate(previews):
        with Image.open(p) as im: thumb=ImageOps.contain(im.convert('RGB'),(790,445))
        x=20+(i%2)*820; y=25+(i//2)*480; canvas.paste(thumb,(x,y)); draw.text((x,y+450),f'{i+1:02}  '+SLIDES[i]['title'].replace('\n',' '),fill='#12332D')
    canvas.save(OUT/f'{NAME}-contact-sheet.png')
    sections=[]
    for i,s in enumerate(SLIDES,1):
        sections.append(f'<section id="slide-{i}" aria-label="Slide {i}"><img src="slides/slide-{i:02}.png" alt="{escape(s["title"])}"><details><summary>Narration and evidence for slide {i}</summary><h2>{escape(s["title"])}</h2><p><b>Say:</b> {escape(s["speak"])}</p><p class="evidence"><b>If asked:</b> {escape(s["evidence"])}</p>'+''.join(f'<a href="{escape(SOURCES[k])}">{escape(k)}</a> ' for k in s['sources'])+'</details></section>')
    html='''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Savia — final pitch</title><style>*{box-sizing:border-box}body{margin:0;background:#0c211d;color:#f8f8f2;font:17px/1.5 Arial,sans-serif}nav{position:sticky;top:0;background:#12332df2;padding:12px 5%;display:flex;gap:16px;align-items:center;z-index:2}a{color:#d9edb7}nav b{margin-right:auto}section{max-width:1280px;margin:24px auto;scroll-margin-top:80px}section img{display:block;width:100%;height:auto}details{padding:16px 24px;background:#214a3f}.evidence{font-size:14px;color:#c6d3ca}summary{cursor:pointer}h2{font-size:24px}.hint{max-width:1280px;margin:16px auto;color:#c6d3ca;font-size:14px}@media print{nav,details,.hint{display:none}section{margin:0;break-after:page}}@media(max-width:700px){nav{font-size:13px;gap:10px}section{margin:12px auto}.hint{padding:12px}}</style><nav><b>SAVIA / FINAL PITCH</b><a href="savia-final-pitch.pptx">Editable PPTX</a><a href="savia-final-pitch.pdf">PDF</a><a href="https://savia-rc-2026.fly.dev">Demo</a></nav><p class="hint">Six slides · English pitch · arrow keys move between slides · evidence and narration beneath each slide.</p>'''+''.join(sections)+'''<script>document.addEventListener('keydown',e=>{if(!['ArrowRight','ArrowLeft'].includes(e.key))return;e.preventDefault();const a=[...document.querySelectorAll('section')];let i=a.reduce((best,s,j)=>Math.abs(s.getBoundingClientRect().top-80)<Math.abs(a[best].getBoundingClientRect().top-80)?j:best,0);i=Math.max(0,Math.min(a.length-1,i+(e.key==='ArrowRight'?1:-1)));a[i].scrollIntoView({behavior:'smooth'});});</script></html>'''
    (OUT/f'{NAME}.html').write_text(html,encoding='utf-8')
    manifest={'schema':'savia-pitch-export/v1','slide_count':6,'files':{p.relative_to(OUT).as_posix():{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(OUT.rglob('*')) if p.is_file() and p.name!='savia-final-pitch-manifest.json'}}
    (OUT/f'{NAME}-manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
    print('Preview, HTML and SHA-256 export manifest complete.')

if __name__=='__main__':
    import sys
    finish() if '--finish' in sys.argv else build()
