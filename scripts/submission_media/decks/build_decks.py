from pathlib import Path
from html import escape
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from shutil import copyfile
from PIL import Image

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "docs" / "submission" / "media" / "decks"
OUT.mkdir(parents=True, exist_ok=True)
ANSWER_SCREENSHOT = ROOT / "docs" / "submission" / "measurements" / "story-es-v2" / "04-actual-answer.png"
FOLLOWUP_SCREENSHOT = ROOT / "docs" / "submission" / "measurements" / "story-es-v2" / "10-rechecked-next-step.png"
FINAL_ANSWER_SCREENSHOT = ROOT / "docs" / "submission" / "measurements" / "team-story-final" / "04-grounded-answer.png"
FINAL_TEAM_SCREENSHOT = ROOT / "docs" / "submission" / "measurements" / "team-story-final" / "11-saved-useful-result.png"
INTEGRATED_CLOSURE_SCREENSHOT = ROOT / "docs" / "submission" / "measurements" / "intended-savia-native" / "completed-inquiry.png"
HERO_SOURCE = ROOT / "docs" / "video" / "pilot" / "keyframe-screen-check.jpg"
ANSWER_CROP = OUT / "savia-answer-crop.png"
FOLLOWUP_CROP = OUT / "savia-followup-crop.png"
FINAL_ANSWER_CROP = OUT / "savia-grounded-answer-crop.png"
FINAL_TEAM_CROP = OUT / "savia-team-saved-result-crop.png"
INTEGRATED_CLOSURE_CROP = OUT / "savia-integrated-closure-crop.png"
HERO_CROP = OUT / "savia-hero-customer.png"
with Image.open(ANSWER_SCREENSHOT) as capture:
    capture.crop((330, 0, 950, 700)).save(ANSWER_CROP)
with Image.open(FOLLOWUP_SCREENSHOT) as capture:
    capture.crop((330, 0, 950, 700)).save(FOLLOWUP_CROP)
with Image.open(FINAL_ANSWER_SCREENSHOT) as capture:
    capture.crop((320, 0, 960, 720)).save(FINAL_ANSWER_CROP)
with Image.open(FINAL_TEAM_SCREENSHOT) as capture:
    capture.crop((320, 0, 960, 720)).save(FINAL_TEAM_CROP)
with Image.open(INTEGRATED_CLOSURE_SCREENSHOT) as capture:
    capture.crop((300, 0, 1060, 860)).save(INTEGRATED_CLOSURE_CROP)
with Image.open(HERO_SOURCE) as capture:
    capture.crop((0, 0, 960, 535)).save(HERO_CROP)

NAVY = "173C35"
PANEL = "244E40"
PANEL2 = "315F4E"
TEAL = "D9EDB7"
CREAM = "F8F9F6"
MUTED = "C8D1C4"
GOLD = "FFC879"
CORAL = "FF907A"
FONT = "Aptos"

JUDGING = [
    {
        "kicker": "FACTORED AI & DATA HACKATHON 2026   /   SAVIA",
        "title": "Ask once.\nCarry on.",
        "sub": "You shouldn’t have to repeat your story or chase updates. Savia keeps your place.",
        "kind": "hero",
        "cards": [("ASK", "Use your own words", "Start with the question on your mind."),
                  ("TEAM", "Find helpful angles", "Evidence and next steps come together."),
                  ("RETURN", "Come back to progress", "Pick up the same inquiry later.")],
        "banner": "MEET SAVIA  ·  YOUR FRIENDLY VOICE ASSISTANT",
        "notes": "Open with the product promise: ask once, let Savia gather useful perspectives, then return to progress without starting over. The home scene is an illustrative fictional TV-spot image, not a service screenshot. The captured evidence that follows uses fictional customer data and informational suggestions."
    },
    {
        "kicker": "02   /   A TEAM IN YOUR CORNER",
        "title": "A grounded answer, in the moment.",
        "sub": "Savia explains the selected charge and the existing record, in the customer's words.",
        "kind": "journey",
        "screen": "grounded-answer",
        "cards": [("MODEL ANSWER", "Useful in 9.338 s", "Grounded to the selected MXN 4,280.75 charge and existing claim."),
                  ("STATUS", "Already in review", "Savia explains the existing record; it creates no new claim."),
                  ("SEPARATE MCP", "2 workers · 2.783 s", "Standalone stdio run; separate from this browser answer.")],
        "caption": "Actual Spanish browser answer · selected fictional charge",
        "banner": "Existing simulated claim · no bank decision or new request",
        "notes": "The screenshot is the actual grounded Spanish answer in docs/submission/measurements/team-story-final/04-grounded-answer.png. The selected model answer was useful and took 9.338 seconds. The browser showed the existing fictional claim and its current review status; it did not create another claim. The informational team interaction completed in 4.278 seconds. Keep these measurements distinct: the separate standalone stdio MCP smoke in docs/submission/assistant/mcp-smoke.json completed two workers in 2.783 seconds; it is not the timing for this browser answer or team interaction. No bank decision or real action is claimed."
    },
    {
        "kicker": "03   /   COME BACK WHEN IT SUITS YOU",
        "title": "Your team’s suggestions stay with you.",
        "sub": "Mark an answer helpful. Start a new chat. Return to the same saved perspectives.",
        "kind": "followthrough",
        "screen": "team-saved-result",
        "cards": [("CUSTOMER CHOICE", "Closed as helpful", "The customer explicitly confirmed the informational answer helped."),
                  ("SAVED TOGETHER", "Two team suggestions", "Both remained available after a new chat and reload."),
                  ("BANK RECORD", "Existing claim", "No new claim was created by the assistant team.")],
        "caption": "Actual Spanish browser capture · closure and saved result",
        "banner": "Informational help saved · existing claim only · no bank decision",
        "notes": "The screenshot is the actual saved-result capture docs/submission/measurements/team-story-final/11-saved-useful-result.png. The customer explicitly closed the informational inquiry as helpful. Both real Savia-team suggestions remained visible after starting a new chat and reloading. The selected bank claim shown in context was already present; the team did not create a new claim or resolve it. The actual UI completion for the team inquiry was 4.278 seconds. The selected grounded answer was separately measured at 9.338 seconds; the standalone MCP smoke took 2.783 seconds. These are different runs and timings."
    },
    {
        "kicker": "04   /   VOICE THAT STAYS WITH YOU",
        "title": "A voice that keeps your place.",
        "sub": "A calm moment, then a spoken confirmation when your explanation helps.",
        "kind": "voice",
        "voice_advice": "Despacio, sí, bien tranquilo, aquí estoy.",
        "voice_result": "Cerró la consulta, ya te ayudó la explicación.",
        "banner": "Actual integrated Savia voice · customer-marked-helpful inquiry · saved context restored",
        "notes": "This is the accepted integrated Savia UI capture from source 9d77a7599128b668b0e34f9c2937eb40b6bd3824, image sha256 484fe8edc07ac37dd78f61deb7a08254352b8ce597fa8265f917706e93d199a5, and served JS hash f06ecbad1ea5089a227a04ea7468aeb8f4b1ef023ac04a98e6b71379771cc8bd. The foreground provider transcript was “Despacio, sí, bien tranquilo, aquí estoy.” (6.6 s; exact 158,400 samples) and received one complete HTTP 200 playback acknowledgement at 21.473 s. While it played, the existing third informational inquiry was marked helpful at 14.949 s; its canonical closure became available at 15.836 s and was queued until the next native request at 21.503 s. The closure transcript was “Cerró la consulta, ya te ayudó la explicación.” (3.95 s; exact 94,800 samples) with one complete HTTP 200 acknowledgement at 26.659 s. The screenshot shows the customer-marked helpful closure and saved team suggestions. This closed an informational case only; no bank decision or refund occurred. The existing case, two completed worker suggestions and prior bank reply were restored; this continuation created zero bank chats, inquiries or workers. It does not claim a new team task or recommendations spoken aloud. Input used typed Spanish and a file-backed fictional WAV through the browser voice control; physical-microphone quality is not qualified, and the raw final ASR contained errors. Optional recorder cleanup diagnostics later exceeded the 180-second cap while awaiting unused response bodies; both exact playback receipts and required product checks had already passed. The diagnostic exit status is not a product failure. Evidence: docs/submission/measurements/intended-savia-native/receipt.json."
    },
    {
        "kicker": "05   /   HOW SAVIA HOLDS THE THREAD",
        "title": "One guide, with clear boundaries.",
        "sub": "Savia keeps the conversation, brings perspectives together, and returns useful progress over time.",
        "kind": "architecture",
        "steps": [("01", "You ask", "Voice or text"), ("02", "Savia host", "Holds context"), ("03", "Savia team", "Two perspectives"), ("04", "Informational read", "Can fail safely"), ("05", "Check back", "30 min · up to 7 days")],
        "banner": "Fictional information service · no bank actions or real resolution",
        "credits": "Prompt flow & decision design: Gloria Yanta Salc     •     Data pipeline & lookup: Carlos Diaz     •     Integration: Savia team",
        "notes": "Explain the current structure without promising production integration. The canonical browser path uses the Savia app host loop to coordinate two direct-provider informational tasks, hold the customer's inquiry, and return progress; the actual team smoke used direct OpenRouter Gemini and both tasks completed in 2,783 ms. The app-owned inquiry policy supports a first reminder after 30 minutes and a tracking window up to seven days; the seven-day operation has not been observed live. Separately, one real-clock simulated receipt check repeated after 30 minutes and deduplicated the unchanged result. The optional generic FLUJO/MCP schedule was qualified with one connected tool call against empty informational state: it completed without queueing a case or making a paid model call, and its 30-minute schedule is installed but disabled. That optional tick is not the browser's Savia team path, and the empty-state run does not demonstrate a scheduled case event. Informational reads may fail; the voice surface reported the read failure. No bank action, real dispute decision or refund is established here. Credits: Gloria Yanta Salc for prompt flow/R0–R18 design, Carlos Diaz for data pipeline and lookup, Savia team for integration. Evidence: docs/submission/runtime-mcp.json and docs/submission/measurements/actual-half-hour-followup.json."
    }
]

PITCH = [
    {
        "kicker": "FACTORED AI & DATA HACKATHON 2026",
        "title": "Ask once.\nCarry on.",
        "sub": "You shouldn’t have to repeat your story or chase updates. Savia keeps your place.",
        "kind": "pitch-hero",
        "banner": "SAVIA  ·  YOUR FRIENDLY VOICE ASSISTANT",
        "notes": "Meet Savia: ask once, carry on. The home scene is an illustrative fictional customer image for the TV spot, not a service screenshot. Savia's product promise is to keep the question moving, bring useful perspectives back, and remember the thread."
    },
    {
        "kicker": "THE TEAM BEHIND THE ANSWER",
        "title": "A grounded answer, in the moment.",
        "sub": "Savia explained the selected charge and existing claim in a useful Spanish answer.",
        "kind": "journey",
        "screen": "grounded-answer",
        "cards": [("MODEL ANSWER", "Useful in 9.338 s", "Grounded to the selected MXN 4,280.75 charge and existing claim."), ("STATUS", "Already in review", "No duplicate claim was created."), ("SEPARATE MCP", "2 workers · 2.783 s", "Standalone stdio run, separate from this browser answer.")],
        "caption": "Actual Spanish browser answer · selected fictional charge",
        "banner": "Existing simulated claim · no bank decision or new request",
        "notes": "Actual grounded answer captured in docs/submission/measurements/team-story-final/04-grounded-answer.png. The selected model answer was useful and took 9.338 seconds. It described the existing fictional claim and current review status; it did not create another claim. The informational team interaction completed in 4.278 seconds. Keep these separate from the standalone two-worker stdio MCP smoke, which completed in 2.783 seconds. No real bank decision or action is claimed."
    },
    {
        "kicker": "COME BACK WHEN IT SUITS YOU",
        "title": "Your team’s suggestions stay with you.",
        "sub": "Confirm that it helped, begin a new chat, and return to the saved perspectives.",
        "kind": "followthrough",
        "screen": "team-saved-result",
        "cards": [("CUSTOMER CHOICE", "Marked helpful", "The customer explicitly closed the informational inquiry."), ("PERSISTENCE", "Two suggestions saved", "They remained after a new chat and reload."), ("BANK RECORD", "Existing claim", "The team did not create a new claim.")],
        "caption": "Actual Spanish browser capture · saved result after return",
        "banner": "Informational help saved · existing claim only · no bank decision",
        "notes": "Actual saved-result capture in docs/submission/measurements/team-story-final/11-saved-useful-result.png. The customer explicitly confirmed the answer was helpful and closed the informational inquiry. Both Savia-team suggestions stayed visible after a new chat and reload. The bank claim was an existing simulated record and was not newly created. The UI team interaction completed in 4.278 seconds; the selected grounded model answer took 9.338 seconds. The 2.783-second MCP smoke was a separate standalone run."
    },
    {
        "kicker": "A VOICE THAT STAYS PRESENT",
        "title": "A voice that keeps your place.",
        "sub": "A calm moment, then a spoken confirmation when your explanation helps.",
        "kind": "voice",
        "voice_advice": "Despacio, sí, bien tranquilo, aquí estoy.",
        "voice_result": "Cerró la consulta, ya te ayudó la explicación.",
        "banner": "Actual integrated Savia voice · customer-marked-helpful inquiry · saved context restored",
        "notes": "Accepted actual integrated Savia UI capture, source 9d77a7599128b668b0e34f9c2937eb40b6bd3824, image sha256 484fe8edc07ac37dd78f61deb7a08254352b8ce597fa8265f917706e93d199a5, served JS hash f06ecbad1ea5089a227a04ea7468aeb8f4b1ef023ac04a98e6b71379771cc8bd. Foreground Spanish transcript: “Despacio, sí, bien tranquilo, aquí estoy.” (6.6 s; exact 158,400 samples), full browser playback acknowledged once with HTTP 200 at 21.473 s. The existing third informational inquiry was marked helpful at 14.949 s during the foreground playback. Its canonical closure was available at 15.836 s and queued until native request 21.503 s. Closure transcript: “Cerró la consulta, ya te ayudó la explicación.” (3.95 s; exact 94,800 samples), acknowledged once with HTTP 200 at 26.659 s. This is an informational helpful closure, not a bank decision or refund. Existing case context, two completed team suggestions and prior bank reply were restored; the continuation created zero bank chats, inquiries or workers. No additional worker concurrency or spoken recommendations are claimed. Input used typed Spanish plus a file-backed fictional WAV through the browser control; the physical microphone is unqualified and raw final ASR contained errors. Optional recorder cleanup exceeded its 180-second cap while awaiting unused response bodies after both playback receipts and required product checks passed; its diagnostic exit 1 is not a product failure. Evidence: docs/submission/measurements/intended-savia-native/receipt.json."
    },
    {
        "kicker": "THE SERVICE UNDER THE SURFACE",
        "title": "Let’s prove value with a bounded pilot.",
        "sub": "Partner with a bank to test the customer-support journey and its handoffs.",
        "kind": "architecture",
        "steps": [("01", "You ask", "Voice or text"), ("02", "Savia host", "Keeps context"), ("03", "Savia team", "Two perspectives"), ("04", "Info read", "May fail safely"), ("05", "Check back", "30 min · up to 7 days")],
        "banner": "BANK PARTNER INVITED · bounded customer-support pilot · no production actions",
        "credits": "Gloria Yanta Salc · prompt flow     |     Carlos Diaz · data pipeline     |     Savia team · integration",
        "pilot_metrics": "Measure: repeat contacts · time to first useful answer · helpful closure · handoff quality",
        "notes": "Close with the invitation for a bank partner to run a bounded customer-support pilot. The value hypothesis is fewer repeat contacts and clearer context for support; this deck claims no ROI or measured reduction. Agree measures for repeat contacts, time to first useful answer, customer helpful closure and handoff quality. Use a fictional, informational scope and no production actions. The canonical browser path uses the Savia app host loop for two direct-provider tasks; it is separate from optional generic FLUJO/MCP scheduling. The latter's qualification made one connected call against empty informational state, queued no case, made no paid model call, and left its installed 30-minute schedule disabled. The 2.783-second standalone MCP smoke is also separate from the browser answer and team UI times. Evidence: docs/submission/runtime-mcp.json."
    }
]

def rgb(hexv):
    return RGBColor.from_string(hexv)

def set_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(NAVY)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(.11))
    bar.fill.solid(); bar.fill.fore_color.rgb = rgb(TEAL); bar.line.fill.background()

def text(slide, x, y, w, h, value, size, color=CREAM, bold=False, font=FONT,
         align=PP_ALIGN.LEFT, margin=0.02, valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(0)
        r = p.add_run(); r.text = line
        r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color)
    return box

def card(slide, x, y, w, h, eyebrow, title, body, accent=TEAL, body_size=15):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(PANEL); shape.line.color.rgb = rgb(PANEL2); shape.line.width = Pt(1)
    text(slide, x+.20, y+.14, w-.40, .28, eyebrow, 10, accent, True)
    text(slide, x+.20, y+.49, w-.40, .52, title, 19, CREAM, True, valign=MSO_ANCHOR.TOP)
    text(slide, x+.20, y+1.06, w-.40, h-1.18, body, body_size, MUTED, valign=MSO_ANCHOR.TOP)

def compact_card(slide,x,y,w,h,eyebrow,title,body,accent=TEAL):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=rgb(PANEL);shape.line.color.rgb=rgb(PANEL2)
    text(slide,x+.18,y+.06,1.62,.22,eyebrow,8,accent,True)
    text(slide,x+1.80,y+.05,2.05,.28,title,12,CREAM,True)
    text(slide,x+3.88,y+.05,w-4.04,h-.10,body,10,MUTED)

def header(slide, spec, idx, total):
    set_bg(slide)
    text(slide, .56, .30, 12.2, .26, spec["kicker"], 10, TEAL, True)
    hero = spec["kind"] in ("hero", "pitch-hero")
    two = "\n" in spec["title"]
    if not hero:
        text(slide, .56, .77, 12.2, 1.14 if two else .72, spec["title"], 30, CREAM, True, valign=MSO_ANCHOR.TOP)
    if spec.get("sub") and not hero:
        text(slide, .58, 2.00 if two else 1.60, 12.15, .42, spec["sub"], 14, MUTED, valign=MSO_ANCHOR.TOP)
    text(slide, 11.5, 7.13, 1.28, .18, f"SAVIA   •   {idx:02}/{total:02}", 8, MUTED, align=PP_ALIGN.RIGHT)

def banner(slide, value, y=6.60, color=PANEL2, font_size=12):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.58), Inches(y), Inches(12.17), Inches(.42))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(PANEL2); sh.line.fill.background()
    text(slide,.78,y+.035,11.78,.32,value,font_size,CREAM,True)

def draw(spec, prs, idx, total):
    slide=prs.slides.add_slide(prs.slide_layouts[6]); header(slide,spec,idx,total); k=spec["kind"]
    if k=="hero":
        text(slide,.68,1.08,5.75,1.26,spec["title"],38,CREAM,True,valign=MSO_ANCHOR.TOP)
        text(slide,.72,2.48,5.75,1.00,spec["sub"],19,MUTED,True,valign=MSO_ANCHOR.TOP)
        frame=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.82), Inches(1.14), Inches(5.90), Inches(3.36))
        frame.fill.solid();frame.fill.fore_color.rgb=rgb(CREAM);frame.line.color.rgb=rgb(PANEL2)
        slide.shapes.add_picture(str(HERO_CROP), Inches(6.86), Inches(1.18), width=Inches(5.82), height=Inches(3.28))
        for i,(label,title,body) in enumerate(spec["cards"]):
            y=3.78+i*.66
            text(slide,.74,y,1.40,.22,f"0{i+1}  {label}",8,TEAL,True)
            text(slide,2.12,y-.01,4.45,.35,title,14,CREAM,True)
        banner(slide,spec["banner"],6.08,GOLD,11)
    elif k=="data":
        x=.58;y=2.25
        panel=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(4.10), Inches(3.42));panel.fill.solid();panel.fill.fore_color.rgb=rgb(PANEL);panel.line.fill.background()
        text(slide,x+.30,y+.28,3.5,.75,spec["stat"][0],37,TEAL,True)
        text(slide,x+.32,y+1.10,3.50,.40,spec["stat"][1],16,CREAM,True)
        text(slide,x+.32,y+1.62,3.50,.58,spec["stat"][2],15,MUTED)
        text(slide,x+.32,y+2.55,3.50,.40,spec["stat"][3],10,GOLD,True)
        for i,(e,t,b) in enumerate(spec["cards"]): compact_card(slide,4.93,2.25+i*1.17,7.82,1.02,e,t,b,[CORAL,CORAL,TEAL][i])
        banner(slide,spec["banner"],6.12)
    elif k in ("journey","followthrough","pitch-experience","pitch-followthrough"):
        frame=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.58), Inches(2.49), Inches(3.58), Inches(3.97))
        frame.fill.solid();frame.fill.fore_color.rgb=rgb(CREAM);frame.line.color.rgb=rgb(PANEL2)
        shot={"grounded-answer":FINAL_ANSWER_CROP,"team-saved-result":FINAL_TEAM_CROP}.get(spec.get("screen"), ANSWER_CROP if k in ("journey","pitch-experience") else FOLLOWUP_CROP)
        slide.shapes.add_picture(str(shot), Inches(.66), Inches(2.55), width=Inches(3.42), height=Inches(3.85))
        if spec.get("cards"):
            colors=[TEAL,GOLD,CORAL]
            for i,(eyebrow,title,body) in enumerate(spec["cards"]):
                compact_card(slide,4.48,2.55+i*1.21,8.26,1.03,eyebrow,title,body,colors[i%len(colors)])
            text(slide,4.52,6.11,8.10,.32,spec.get("caption","Actual Spanish browser capture"),10,MUTED)
        else:
            card(slide,4.48,2.55,8.26,3.90,"WHAT THE CUSTOMER GETS","Clarity without false certainty","• One customer-owned charge\n• A grounded explanation\n• A clear next step in Spanish or Portuguese\n• Follow-up after opt-in and receipt read-back",TEAL,15)
        banner(slide,spec["banner"],6.58,GOLD if k in ("journey","followthrough") else PANEL2,9)
    elif k=="architecture":
        for i,(n,t,b) in enumerate(spec["steps"]):
            x=.58+i*2.48
            s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.55), Inches(2.28), Inches(1.80));s.fill.solid();s.fill.fore_color.rgb=rgb(PANEL);s.line.color.rgb=rgb(PANEL2)
            text(slide,x+.16,2.70,1.8,.25,n,10,TEAL,True)
            text(slide,x+.16,3.05,2.0,.52,t,16,CREAM,True)
            text(slide,x+.16,3.62,2.0,.50,b,11,MUTED)
            if i<4: text(slide,x+2.27,3.18,.23,.4,"→",17,TEAL,True,align=PP_ALIGN.CENTER)
        banner(slide,spec["banner"],4.82,GOLD,10)
        text(slide,.62,5.43,12.0,.36,spec["credits"],11,CREAM,True)
        if spec.get("pilot_metrics"):
            text(slide,.62,5.88,12.0,.38,spec["pilot_metrics"],10,MUTED)
    elif k=="voice":
        frame=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.58), Inches(2.49), Inches(3.58), Inches(3.97))
        frame.fill.solid();frame.fill.fore_color.rgb=rgb(CREAM);frame.line.color.rgb=rgb(PANEL2)
        slide.shapes.add_picture(str(INTEGRATED_CLOSURE_CROP), Inches(.66), Inches(2.55), width=Inches(3.42), height=Inches(3.85))
        card(slide,4.48,2.55,8.26,1.47,"FOREGROUND · 6.6 S · FULL ACK 200","Calm reassurance",f'“{spec["voice_advice"]}”',TEAL,15)
        card(slide,4.48,4.22,8.26,1.47,"HELPFUL CLOSURE · 3.95 S · FULL ACK 200","Customer-confirmed closure",f'“{spec["voice_result"]}”',GOLD,15)
        banner(slide,spec["banner"],6.62,PANEL2,10)
    elif k in ("evidence","pitch-limits"):
        for i,(e,t,b) in enumerate(spec["cards"]): card(slide,.58+i*4.13,2.60,3.88,2.90,e,t,b,[TEAL,GOLD,CORAL][i],15)
        banner(slide,spec["banner"],6.02,GOLD,11)
    elif k=="bank-value":
        for i,(e,t,b) in enumerate(spec["cards"]): card(slide,.58+i*4.13,2.67,3.88,2.84,e,t,b,[TEAL,GOLD,CORAL][i],15)
        banner(slide,spec["banner"],5.91,GOLD,11)
    elif k=="pitch-value":
        for i,(e,t,b) in enumerate(spec["cards"]): card(slide,.58+i*4.13,2.67,3.88,2.84,e,t,b,[TEAL,GOLD,CORAL][i],15)
        banner(slide,spec["banner"],5.91,GOLD,11)
    elif k=="pitch-hero":
        text(slide,.68,1.16,5.75,1.48,spec["title"],42,CREAM,True,valign=MSO_ANCHOR.TOP)
        text(slide,.72,2.80,5.65,1.00,spec["sub"],21,TEAL,True,valign=MSO_ANCHOR.TOP)
        frame=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.82), Inches(1.14), Inches(5.90), Inches(3.36))
        frame.fill.solid();frame.fill.fore_color.rgb=rgb(CREAM);frame.line.color.rgb=rgb(PANEL2)
        slide.shapes.add_picture(str(HERO_CROP), Inches(6.86), Inches(1.18), width=Inches(5.82), height=Inches(3.28))
        text(slide,.74,4.05,5.9,.86,"A small team finds useful perspectives. Savia keeps the thread while you get on with your day.",17,MUTED)
        banner(slide,spec["banner"],6.08,PANEL2,11)
    elif k=="pitch-problem":
        text(slide,.72,2.45,4.0,1.0,"12,297",42,TEAL,True)
        text(slide,.75,3.46,3.8,.42,"of 67,095 complaints",17,CREAM,True)
        text(slide,.75,4.06,3.85,.52,"18.3% tagged as unrecognized charge",15,MUTED)
        card(slide,5.0,2.50,7.75,2.95,"THE DATA CHECK", "Bad links change the design", "Complaint-to-interaction links are empty. Populated complaint-to-product links fail owner matching. Savia resolves the charge from the signed-in customer's own records.",CORAL,16)
        banner(slide,spec["banner"],6.10,GOLD,12)
    elif k=="pitch-design":
        text(slide,.72,2.45,11.9,1.05,"Portal   →   trusted Savia host   →   bounded language service   →   in-process Banking MCP",21,CREAM,True)
        card(slide,.72,3.88,5.77,1.55,"LANGUAGE", "Model supports the conversation", "The captured reply was a deterministic host fallback, not a model success.",TEAL,13)
        card(slide,6.80,3.88,5.77,1.55,"AUTHORITY", "Services retain control", "Owner context, consent, scoped reads and simulated receipt checks.",GOLD,13)
        banner(slide,spec["banner"],6.10)
    elif k=="pitch-pilot":
        for i,(e,t,b) in enumerate(spec["cards"]): card(slide,.58+i*4.13,2.60,3.88,2.88,e,t,b,[TEAL,GOLD,CORAL][i],15)
        banner(slide,spec["banner"],5.98,GOLD,11)
    return slide

def make_deck(specs, name, title, tagline):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    prs.core_properties.title=title; prs.core_properties.subject=tagline; prs.core_properties.author="Savia team"
    total=len(specs)
    for idx,spec in enumerate(specs,1):
        slide=draw(spec,prs,idx,total)
        # Speaker notes are exported as a parallel editable Markdown script below;
        # PowerPoint note pages are added by the native Office export step.
    path=OUT/f"{name}.pptx"; prs.save(path)
    write_html(specs,name,title)
    write_notes(specs,name,title)

def write_notes(specs,name,title):
    lines=[f"# {title} — speaker notes", "", "All customer examples are fictional. Source-data counts refer to the organizer-supplied synthetic snapshot; no live bank action is claimed.", ""]
    for i,s in enumerate(specs,1): lines += [f"## Slide {i}: {s['title'].replace(chr(10),' ')}", "", s["notes"], ""]
    (OUT/f"{name}-speaker-notes.md").write_text("\n".join(lines),encoding="utf-8")
    import json
    (OUT/f"{name}-speaker-notes.json").write_text(json.dumps([s["notes"] for s in specs],ensure_ascii=False,indent=2),encoding="utf-8")

def write_html(specs,name,title):
    slides=[]
    for i,s in enumerate(specs,1):
        cards=s.get("cards",[])
        card_html="".join(f'<article><small>{escape(a)}</small><h2>{escape(b)}</h2><p>{escape(c)}</p></article>' for a,b,c in cards)
        stat=s.get("stat")
        data=(f'<div class="stat"><strong>{escape(stat[0])}</strong><b>{escape(stat[1])}</b><p>{escape(stat[2])}</p><small>{escape(stat[3])}</small></div>' if stat else '')
        steps=s.get("steps",[])
        step_html=''.join(f'<article><small>{n}</small><h2>{escape(t)}</h2><p>{escape(b)}</p></article>' for n,t,b in steps)
        content=f'<div class="cards">{card_html}</div>{data}<div class="steps">{step_html}</div>'
        if s.get("kind") in ("hero","pitch-hero"):
            image='<img src="savia-hero-customer.png" alt="Illustrative fictional Savia customer scene">'
            hero_title=f'<h2>{escape(s["title"]).replace(chr(10),"<br>")}</h2><p>{escape(s["sub"])}</p>'
            small=("".join(f'<li>{escape(t)}</li>' for _,t,_ in cards) if cards else '<li>Talk it through</li><li>Get a clear next step</li><li>Come back without starting over</li>')
            content=f'<div class="hero-show"><div>{hero_title}<ul>{small}</ul></div>{image}</div>'
        if s.get("kind") == "voice":
            content=(f'<div class="experience"><img src="savia-integrated-closure-crop.png" alt="Integrated Savia interface showing the completed helpful inquiry and saved suggestions">'
                     f'<div class="experience-copy"><h2>Foreground · 6.6 seconds</h2><p>“{escape(s["voice_advice"])}”</p>'
                     f'<h2>Helpful closure · 3.95 seconds</h2><p>“{escape(s["voice_result"])}”</p>'
                     '<p>Informational closure only; no bank decision or refund.</p></div></div>')
        if s.get("kind") in ("journey","followthrough","pitch-experience","pitch-followthrough"):
            screen=s.get("screen")
            details=''.join(f'<article><small>{escape(a)}</small><h2>{escape(b)}</h2><p>{escape(c)}</p></article>' for a,b,c in cards)
            if screen=="grounded-answer":
                shot="savia-grounded-answer-crop.png"; title="A grounded answer, in the moment."; story="The actual Spanish model answer explained the selected charge and its existing simulated claim."
            elif screen=="team-saved-result":
                shot="savia-team-saved-result-crop.png"; title="Your team's suggestions stay with you."; story="The customer marked the informational result helpful; both suggestions remained after a new chat and reload."
            else:
                is_follow=s.get("kind") in ("followthrough","pitch-followthrough")
                story=("Actual Spanish browser capture. Explicit consent was followed by a verified simulated receipt and an opt-in status recheck after reload. No bank decision or refund was verified." if is_follow else "Actual Spanish browser capture. The assistant displayed recorded transaction facts and explained that a chat reply alone does not submit a request.")
                shot="savia-followup-crop.png" if is_follow else "savia-answer-crop.png"
                title="The receipt is still registered" if is_follow else "The answer shows what is known"
            content=f'<div class="experience"><img src="{shot}" alt="Actual fictional Savia customer conversation capture"><div class="experience-copy"><h2>{escape(title)}</h2><p>{escape(story)}</p><div class="capture-cards">{details}</div></div></div>'
        if s.get("credits"): content+=f'<p class="credits">{escape(s["credits"])}</p>'
        if s.get("banner"): content+=f'<div class="banner">{escape(s["banner"])}</div>'
        slides.append(f'<section><header>{escape(s["kicker"])}</header><h1>{escape(s["title"]).replace(chr(10),"<br>")}</h1><p class="sub">{escape(s.get("sub",""))}</p>{content}<footer>SAVIA　•　{i:02}/{len(specs):02}</footer></section>')
    html='''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'''+escape(title)+'''</title><style>
*{box-sizing:border-box}body{margin:0;background:#e7eeea;color:#f8f9f6;font:16px/1.45 Aptos,Inter,Arial,sans-serif}section{position:relative;width:min(100vw,1440px);min-height:min(56.25vw,810px);margin:24px auto;padding:54px 6%;background:#173c35;border-top:10px solid #d9edb7;overflow:hidden}header{color:#d9edb7;font-size:13px;font-weight:700;letter-spacing:.1em}h1{font-size:clamp(32px,4vw,58px);line-height:1.04;margin:32px 0 16px}.sub{color:#c8d1c4;max-width:980px}.cards{display:flex;gap:20px;margin-top:56px}article{flex:1;min-height:190px;padding:22px;border:1px solid #315f4e;border-radius:18px;background:#244e40}article small{color:#d9edb7;font-weight:700;letter-spacing:.06em}article h2{font-size:22px;line-height:1.12;margin:18px 0 12px}article p{color:#c8d1c4;margin:0}.stat{position:absolute;left:7%;top:42%;width:28%;}.stat strong{display:block;color:#d9edb7;font-size:58px}.stat b{display:block}.stat p{color:#c8d1c4}.steps{display:flex;gap:14px;margin-top:50px}.steps article{min-height:140px}.steps article h2{font-size:19px}.banner{margin-top:28px;padding:13px 18px;border-radius:8px;background:#315f4e;color:#f8f9f6;font-weight:700;font-size:13px}.credits{margin-top:26px;color:#f8f9f6;font-weight:600}footer{position:absolute;right:6%;bottom:25px;color:#c8d1c4;font-size:11px;letter-spacing:.1em}.experience{display:flex;gap:26px;align-items:flex-start;margin-top:42px}.experience img{width:52%;height:auto;border-radius:8px;border:1px solid #315f4e}.experience-copy{flex:1;background:#244e40;border:1px solid #315f4e;border-radius:18px;padding:24px}.experience-copy h2{margin-top:0;color:#d9edb7}.experience-copy p{color:#c8d1c4} @media(max-width:800px){section{min-height:100vh;padding:28px 5%;margin:0}.cards,.steps,.experience{display:block;margin-top:30px}.experience img{width:100%;margin-bottom:18px}.cards article,.steps article{margin-bottom:12px}.stat{position:static;width:auto}.banner{font-size:11px}footer{position:static;text-align:right;margin-top:24px}}
</style><body>'''+''.join(slides)+'''</body></html>'''
    html=html.replace(" @media(max-width:800px)", " .hero-show{display:grid;grid-template-columns:1fr 1fr;gap:36px;align-items:center;margin-top:36px}.hero-show h2{font-size:42px;line-height:1.06}.hero-show p{color:#c8d1c4;font-size:21px}.hero-show img{width:100%;border-radius:10px;border:1px solid #315f4e}.hero-show ul{padding-left:18px;color:#d9edb7;line-height:2}.capture-cards{display:grid;gap:9px;margin-top:14px}.capture-cards article{min-height:0;padding:10px 12px}.capture-cards article small{font-size:9px}.capture-cards article h2{font-size:14px;margin:4px 0}.capture-cards article p{font-size:11px} @media(max-width:800px)")
    (OUT/f"{name}.html").write_text(html,encoding="utf-8")

make_deck(JUDGING,"savia-submission-deck","Savia — Hackathon Submission","A focused, multilingual customer-service prototype")
make_deck(PITCH,"savia-pitch-deck","Savia — Short Pitch","A clearer next step for an unfamiliar charge")
print(f"Wrote two five-slide decks and HTML/notes exports to {OUT}")
